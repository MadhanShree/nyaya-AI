"""Turn an uploaded file into page-tagged text, with strict type checks."""

import io
import os

from docx import Document
from pypdf import PdfReader

MAX_PAGES = 100


class ExtractionError(ValueError):
    """The message is safe to show to users."""


def extract_pages(filename: str, data: bytes) -> list[tuple[int, str]]:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf" and data[:5] == b"%PDF-":
        pages = _pdf(data)
    elif ext == ".docx" and data[:2] == b"PK":
        pages = [(1, _docx(data))]
    elif ext == ".txt" and b"\x00" not in data:
        pages = [(1, data.decode("utf-8", errors="replace"))]
    else:
        raise ExtractionError("Upload a valid PDF, DOCX or TXT file.")
    pages = [(n, t) for n, t in pages if t.strip()]
    if not pages:
        raise ExtractionError("No readable text found. Scanned documents need OCR first.")
    return pages


def _pdf(data: bytes) -> list[tuple[int, str]]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Password-protected PDFs are not supported.")
        if len(reader.pages) > MAX_PAGES:
            raise ExtractionError(f"Documents are limited to {MAX_PAGES} pages.")
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError("This PDF could not be read.") from e


def _docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
        rows = [" | ".join(c.text.strip() for c in r.cells) for t in doc.tables for r in t.rows]
        return "\n".join([p.text for p in doc.paragraphs] + rows)
    except Exception as e:
        raise ExtractionError("This DOCX could not be read.") from e
