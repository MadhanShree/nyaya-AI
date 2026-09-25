import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openai import OpenAIError
from pypdf import PdfWriter

from nyayaai import main as app_module
from nyayaai.chunk import chunk_pages
from nyayaai.extract import ExtractionError, extract_pages
from nyayaai.llm import GroqClient, LLMError, parse_json
from nyayaai.main import app, get_llm
from nyayaai.pipeline import NOT_FOUND, analyze, answer_question
from nyayaai.retrieve import Retriever

SAMPLE = """RENTAL AGREEMENT
This agreement is made on 1 January 2026 between Ravi Kumar (Landlord) and Meena Devi (Tenant).

1. Rent
The tenant shall pay monthly rent of Rs. 12,000 on or before the 5th of each month.

2. Security Deposit
The tenant has paid a refundable security deposit of Rs. 36,000.

3. Term
This agreement runs for 11 months from 1 January 2026.

4. Termination
Either party may terminate this agreement by giving 30 days written notice.

5. Maintenance
The tenant shall keep the premises in good condition and pay for minor repairs.
"""


class Fake:
    def __init__(self, reply):
        self.reply, self.calls = reply, 0

    def complete(self, system, user):
        self.calls += 1
        return self.reply if isinstance(self.reply, str) else json.dumps(self.reply)


def sample_chunks():
    return chunk_pages([(1, SAMPLE)])


def clause(label):
    return next(c for c in sample_chunks() if c.label == label)


def test_chunks_follow_clauses():
    labels = [c.label for c in sample_chunks()]
    assert labels == ["Preamble", "Clause 1", "Clause 2", "Clause 3", "Clause 4", "Clause 5"]
    assert "30 days" in clause("Clause 4").text


def test_unstructured_text_falls_back_to_pages():
    chunks = chunk_pages([(1, "Just one paragraph.\n\nAnother paragraph.")])
    assert [c.label for c in chunks] == ["Page 1", "Page 1"]


def test_retrieval_finds_the_right_clause():
    r = Retriever(sample_chunks())
    assert r.search("When can I terminate this agreement?")[0][0].label == "Clause 4"
    assert r.search("How much is the deposit?")[0][0].label == "Clause 2"
    assert r.search("how do I end the contract")[0][0].label == "Clause 4"
    assert r.search("What is the capital of France?") == []


def test_analysis_drops_invented_citations():
    reply = {
        "document_type": "Rental agreement",
        "summary": "A rental deal.",
        "parties": [{"name": "Ravi Kumar", "role": "Landlord"}],
        "key_facts": [{"label": "Rent", "detail": "Rs. 12,000", "sources": ["c2", "c99"]}],
        "obligations": [],
        "risks": [{"issue": "Made up", "why": "No clause", "sources": ["c99"]}],
    }
    result = analyze(Fake(reply), sample_chunks())
    assert result["key_facts"][0]["sources"] == ["c2"]
    assert result["risks"] == []


def test_answer_is_cited_and_accepts_fenced_json():
    reply = '```json\n{"answerable": true, "answer": "30 days written notice.", "sources": ["c5", "c99"]}\n```'
    out = answer_question(Fake(reply), Retriever(sample_chunks()), "When can I terminate this agreement?")
    assert out["found"] and [c["label"] for c in out["citations"]] == ["Clause 4"]


def test_no_retrieval_hit_skips_the_model():
    llm = Fake({})
    out = answer_question(llm, Retriever(sample_chunks()), "What is the capital of France?")
    assert out["answer"] == NOT_FOUND and llm.calls == 0


def test_uncited_answer_is_treated_as_not_found():
    llm = Fake({"answerable": True, "answer": "Trust me.", "sources": ["c99"]})
    assert not answer_question(llm, Retriever(sample_chunks()), "When can I terminate?")["found"]


def test_bad_model_output_and_bad_files_are_rejected():
    with pytest.raises(LLMError):
        parse_json("no json here")
    with pytest.raises(ExtractionError):
        extract_pages("virus.exe", b"MZ")
    with pytest.raises(ExtractionError):
        extract_pages("fake.pdf", b"not a pdf")


class Router:
    def complete(self, system, user):
        if "TASK: ANALYZE" in user:
            return json.dumps(
                {
                    "document_type": "Rental agreement",
                    "summary": "A rental deal.",
                    "parties": [],
                    "key_facts": [],
                    "obligations": [],
                    "risks": [{"issue": "Short notice", "why": "30 days only", "sources": ["c5"]}],
                }
            )
        return json.dumps(
            {"answerable": True, "answer": "Either side can end it with 30 days notice.", "sources": ["c5"]}
        )


@pytest.fixture
def client():
    app.dependency_overrides[get_llm] = lambda: Router()
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_upload_then_ask(client):
    r = client.post("/api/documents", files={"file": ("rent.txt", SAMPLE.encode(), "text/plain")})
    assert r.status_code == 200
    body = r.json()
    assert body["analysis"]["risks"][0]["sources"] == ["c5"]
    a = client.post(f"/api/documents/{body['doc_id']}/ask", json={"question": "When can I terminate?"}).json()
    assert a["found"] and a["citations"][0]["label"] == "Clause 4"


def test_api_rejects_bad_uploads_and_unknown_documents(client):
    assert client.post("/api/documents", files={"file": ("a.exe", b"MZ")}).status_code == 422
    big = b"a" * (5 * 1024 * 1024 + 1)
    assert client.post("/api/documents", files={"file": ("a.txt", big)}).status_code == 413
    assert client.post("/api/documents/nope/ask", json={"question": "Anything here?"}).status_code == 404


def test_security_headers_and_page(client):
    r = client.get("/")
    assert r.status_code == 200 and r.headers["X-Content-Type-Options"] == "nosniff"


class Seq:
    def __init__(self, *replies):
        self.replies, self.calls = list(replies), 0

    def complete(self, system, user):
        self.calls += 1
        return self.replies.pop(0)


def test_non_english_question_is_translated_for_retrieval():
    answer = json.dumps({"answerable": True, "answer": "30 दिन का नोटिस।", "sources": ["c5"]})
    llm = Seq("When can I terminate this agreement?", answer)
    out = answer_question(llm, Retriever(sample_chunks()), "मैं यह समझौता कब समाप्त कर सकता हूँ?", language="hi")
    assert out["found"] and llm.calls == 2 and "वकील" in out["disclaimer"]


def test_language_instruction_reaches_the_model():
    seen = []

    class Spy(Fake):
        def complete(self, system, user):
            seen.append(system)
            return super().complete(system, user)

    analyze(Spy({"summary": "x"}), sample_chunks(), "ta")
    assert "Tamil" in seen[0]


def test_docx_tables_are_read():
    doc = Document()
    doc.add_paragraph("Lease")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Rent", "Rs. 9,000"
    buf = io.BytesIO()
    doc.save(buf)
    assert "Rs. 9,000" in extract_pages("lease.docx", buf.getvalue())[0][1]


def test_rate_limit(client, monkeypatch):
    monkeypatch.setattr(app_module, "RATE_LIMIT", 2)
    app_module.HITS.clear()
    codes = [client.post("/api/documents/x/ask", json={"question": "Anything?"}).status_code for _ in range(3)]
    assert codes == [404, 404, 429]
    app_module.HITS.clear()


def test_pdf_without_text_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ExtractionError, match="No readable text"):
        extract_pages("scan.pdf", buf.getvalue())


def test_groq_client_requires_key_and_hides_service_errors(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    with pytest.raises(LLMError):
        GroqClient()
    monkeypatch.setenv("GROQ_API_KEY", "test")
    client = GroqClient()

    def boom(**kwargs):
        raise OpenAIError("secret internal detail")

    monkeypatch.setattr(client._client.chat.completions, "create", boom)
    with pytest.raises(LLMError, match="could not answer") as err:
        client.complete("system", "user")
    assert "secret" not in str(err.value)


class Broken:
    def complete(self, system, user):
        raise LLMError("AI down")


def test_model_failure_becomes_502(client):
    app.dependency_overrides[get_llm] = lambda: Broken()
    r = client.post("/api/documents", files={"file": ("a.txt", b"1. Rent\n2. Term\nrent text")})
    assert r.status_code == 502 and r.json()["detail"] == "AI down"


def test_missing_api_key_is_503(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    app_module._client.cache_clear()
    r = TestClient(app).post("/api/documents", files={"file": ("a.txt", b"hello")})
    assert r.status_code == 503
